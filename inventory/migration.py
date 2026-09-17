"""Read Excel without executing macros. Produce a reviewable, all-or-nothing migration."""
import argparse
from collections import Counter
from datetime import date, datetime
import json
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

from .core import ValidationError, day, digest, integer, normalize_component, normalize_event, ref_id

DEFAULT_PROFILE = {
    "balance": {"sheet": "Base Supply", "header_row": 4,
                "columns": {"ref_id": "REF ID", "implant_type": "Implant Type", "size": "Item / Size",
                            "component_type": "Component", "opening_quantity": "Base Supply"}},
    "thresholds": {"sheet": "Inventory Full Supply", "header_row": 5,
                   "columns": {"ref_id": "REF ID", "reorder_level": "Reorder Quantity"}},
    "log": {"sheet": "Component Change Log", "header_row": 5,
            "columns": {"date": "Date", "ref_id": "REF ID", "quantity": "Quantity", "tooth": "Tooth #",
                        "placement": "Placed / O Box", "patient": "Patient", "periodontist": "Periodontist",
                        "referring_dentist": "Referring Dentist", "assistant": "Assistant", "notes": "Notes"}},
    "component_aliases": {"Body": "Implant Body", "HA": "Healing Abutment", "CS": "Cover Screw",
                          "Prosth (Crown)": "Crown"},
    "manufacturer_prefixes": ["Straumann", "Nobel"],
    "ref_aliases": {},
}


def rows(workbook, spec):
    sheet = workbook[spec["sheet"]]
    iterator = sheet.iter_rows(min_row=spec["header_row"], values_only=True)
    header = [str(v or "").strip() for v in next(iterator)]
    indexes = {}
    for key, label in spec["columns"].items():
        if header.count(label) != 1:
            raise ValidationError(f"Sheet {sheet.title}: header {label!r} must appear exactly once")
        indexes[key] = header.index(label)
    for number, row in enumerate(iterator, spec["header_row"] + 1):
        selected = {key: row[index] if index < len(row) else None for key, index in indexes.items()}
        if any(value is not None for value in selected.values()):
            yield number, selected


def whole(value, minimum=0):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return integer(value, minimum)


def excel_day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return day(value)


def prepare(dashboard_path, log_path, cutoff, profile=None):
    from openpyxl import load_workbook
    profile = profile or DEFAULT_PROFILE
    cutoff = day(cutoff)
    if cutoff >= date.today().isoformat():
        raise ValidationError("Use an end-of-day opening balance dated before today")
    aliases = {ref_id(k): ref_id(v) for k, v in profile.get("ref_aliases", {}).items()}
    def resolve(value):
        identity = ref_id(value)
        return aliases.get(identity, identity)
    issues, warnings, components, events = [], [], {}, []
    ignored = 0
    dashboard = load_workbook(dashboard_path, read_only=True, data_only=True)
    log = None
    try:
        thresholds = {}
        if profile.get("thresholds"):
            for number, row in rows(dashboard, profile["thresholds"]):
                try:
                    identity = resolve(row["ref_id"])
                    if identity in thresholds:
                        raise ValidationError("Duplicate reorder threshold REF ID")
                    thresholds[identity] = whole(row["reorder_level"])
                except ValueError as error:
                    issues.append(f"Threshold row {number}: {error}")
        for number, row in rows(dashboard, profile["balance"]):
            try:
                identity = resolve(row["ref_id"])
                if identity in components:
                    raise ValidationError(f"Duplicate REF ID {identity}")
                implant_type = str(row.get("implant_type") or "").strip()
                manufacturer = str(row.get("manufacturer") or "").strip()
                if not manufacturer:
                    for prefix in profile.get("manufacturer_prefixes", []):
                        if implant_type.casefold().startswith(prefix.casefold() + " "):
                            manufacturer, implant_type = prefix, implant_type[len(prefix):].strip()
                            break
                component_type = str(row.get("component_type") or "")
                component = normalize_component({**row, "ref_id": identity, "implant_type": implant_type,
                    "manufacturer": manufacturer, "reorder_level": thresholds.get(identity, row.get("reorder_level", 0)),
                    "component_type": profile.get("component_aliases", {}).get(component_type, component_type)})
                component["opening_quantity"] = whole(row.get("opening_quantity"))
                components[identity] = component
            except ValueError as error:
                issues.append(f"Balance row {number}: {error}")
        occurrences = Counter()
        if log_path:
            log = load_workbook(log_path, read_only=True, data_only=True)
            for number, row in rows(log, profile["log"]):
                try:
                    if not row.get("ref_id") and row.get("quantity") is None:
                        ignored += 1
                        warnings.append(f"Log row {number}: no REF ID or quantity; annotation row omitted")
                        continue
                    quantity = whole(row.get("quantity"), -1_000_000)
                    if quantity == 0:
                        ignored += 1
                        warnings.append(f"Log row {number}: zero quantity; no stock movement, omitted")
                        continue
                    identity = resolve(row["ref_id"])
                    if identity not in components:
                        raise ValidationError(f"Unknown REF ID {identity}; add an explicit ref_aliases mapping or correct the workbook")
                    placement = str(row.get("placement") or "").strip()
                    if placement.casefold() in ("o box", '"o" box', "o-box"):
                        placement = "O Box"
                    elif placement.casefold() == "placed":
                        placement = "Placed"
                    data = {**row, "ref_id": identity, "date": excel_day(row["date"]),
                            "quantity": abs(quantity), "kind": "usage" if quantity < 0 else "addition",
                            "placement": placement, "signed_by": row.get("signed_by") or row.get("assistant") or ""}
                    if data["date"] > date.today().isoformat():
                        raise ValidationError("Future-dated entry")
                    # Row numbers are not IDs: sorting a workbook must not create new movements.
                    clean = normalize_event({**data, "id": str(uuid5(NAMESPACE_URL, "temporary"))})
                    clean.pop("id")
                    signature = digest(clean)
                    occurrences[signature] += 1
                    clean["id"] = str(uuid5(NAMESPACE_URL, f"implant-inventory:{signature}:{occurrences[signature]}"))
                    if occurrences[signature] > 1:
                        warnings.append(f"Log row {number}: identical entry preserved; review whether it is a duplicate")
                    events.append(clean)
                except (ValueError, TypeError) as error:
                    issues.append(f"Log row {number}: {error}")
    finally:
        dashboard.close()
        if log:
            log.close()
    balances = {identity: c["opening_quantity"] for identity, c in components.items()}
    active = [event for event in events if event["date"] > cutoff]
    for event in active:
        balances[event["ref_id"]] += event["quantity"] * (-1 if event["kind"] == "usage" else 1)
    negatives = {identity: quantity for identity, quantity in balances.items() if quantity < 0}
    if negatives:
        warnings.append(f"{len(negatives)} negative projected balances require a physical count")
    bundle = {"version": 1, "cutoff": cutoff, "components": list(components.values()), "events": events}
    report = {"components": len(components), "entries": len(events), "archived_entries": len(events)-len(active),
              "stock_changing_entries": len(active), "ignored_rows": ignored, "issues": issues, "warnings": warnings,
              "negative_balances": negatives, "projected_total": sum(balances.values()),
              "cutoff": cutoff, "ready": not issues and bool(components)}
    return bundle, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dashboard", required=True)
    parser.add_argument("--log")
    parser.add_argument("--cutoff", required=True, help="Opening balance is at END of this YYYY-MM-DD date")
    parser.add_argument("--profile", help="JSON mapping; see examples/import-profile.json")
    parser.add_argument("--output", default="migration.json")
    args = parser.parse_args()
    profile = json.loads(Path(args.profile).read_text(encoding="utf-8")) if args.profile else None
    bundle, report = prepare(args.dashboard, args.log, args.cutoff, profile)
    report_path = Path(args.output).with_suffix(".report.json")
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    if report["ready"]:
        Path(args.output).write_text(json.dumps(bundle, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"Prepared {args.output}; review {report_path} before importing")
    else:
        print(f"Import blocked: {len(report['issues'])} issues. See {report_path}. No bundle written.")
        raise SystemExit(1)


if __name__ == "__main__":
    main()
